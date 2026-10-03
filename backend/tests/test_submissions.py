import uuid
from datetime import timedelta
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from assignments.models import Submission,Attachment

def report(content=b"Independent report text "*8):
    return SimpleUploadedFile("report.txt",content)
def code(content=b"def total(values):\n    return sum(values)\n"):
    return SimpleUploadedFile("main.py",content)
def upload(client, assignment, request_id=None, **extra):
    data={"request_id":str(request_id or uuid.uuid4()),"report_files":report(),"code_files":code(),**extra}
    return client.post(f"/api/assignments/{assignment.id}/submissions/",data,format="multipart")

def test_report_and_code_required(client, student, assignment):
    client.force_login(student)
    r=client.post(f"/api/assignments/{assignment.id}/submissions/",{"request_id":str(uuid.uuid4()),"report_files":report()},format="multipart")
    assert r.status_code==400
    assert Submission.objects.count()==0

def test_idempotent_retry_single_version(client, student, assignment):
    client.force_login(student); key=uuid.uuid4()
    first=upload(client,assignment,key);second=upload(client,assignment,key)
    assert first.status_code==201
    assert second.status_code in (200,201)
    assert first.json()["id"]==second.json()["id"]
    assert Submission.objects.count()==1

def test_request_id_content_conflict(client, student, assignment):
    client.force_login(student);key=uuid.uuid4()
    assert upload(client,assignment,key).status_code==201
    assert upload(client,assignment,key,comment="changed").status_code==409

def test_late_submission_policy(client, student, assignment):
    client.force_login(student);assignment.deadline=timezone.now()-timedelta(seconds=1);assignment.save()
    assert upload(client,assignment).status_code==400
    assignment.allow_late=True;assignment.save()
    result=upload(client,assignment)
    assert result.status_code==201
    assert result.json()["is_late"] is True

def test_upload_failure_is_atomic(client, student, assignment, monkeypatch, settings):
    from pathlib import Path
    client.force_login(student)
    original=Path.open
    def broken(self,*args,**kwargs):
        if self.is_relative_to(settings.MEDIA_ROOT) and args and args[0] in ("wb","xb"):
            raise OSError("simulated full disk")
        return original(self,*args,**kwargs)
    monkeypatch.setattr(Path,"open",broken)
    response=upload(client,assignment)
    assert response.status_code==500
    assert Submission.objects.count()==0 and Attachment.objects.count()==0
    assert list(settings.MEDIA_ROOT.rglob("*"))==[]

def test_attachment_download_scoped(client, student, other_student, assignment):
    client.force_login(student);created=upload(client,assignment)
    assert created.status_code==201
    ident=created.json()["attachments"][0]["id"]
    client.force_login(other_student)
    assert client.get(f"/api/attachments/{ident}/download/").status_code==404

def test_scanned_pdf_not_zero_similarity(client,student,assignment):
    from io import BytesIO
    from pypdf import PdfWriter
    buf=BytesIO();pdf=PdfWriter();pdf.add_blank_page(width=200,height=200);pdf.write(buf)
    client.force_login(student)
    response=upload(client,assignment,report_files=SimpleUploadedFile("scan.pdf",buf.getvalue()))
    assert response.status_code==201
    assert response.json()["attachments"][0]["extraction_status"]=="unsupported"

def test_resubmit_versions_and_server_ownership(client,student,other_student,assignment):
    client.force_login(student)
    first=upload(client,assignment)
    second=upload(client,assignment,student_id=other_student.id,version=900)
    assert first.status_code==201 and second.status_code==201
    assert second.json()["version"]==2
    assert Submission.objects.get(pk=second.json()["id"]).student_id==student.id

def test_unsupported_and_oversize(client,student,assignment):
    client.force_login(student)
    assert upload(client,assignment,report_files=SimpleUploadedFile("old.doc",b"bad")).status_code==400
    assert upload(client,assignment,report_files=report(b"x"*(20*1024*1024+1))).status_code==413

def test_closed_or_archived_reject_submission(client,student,assignment):
    client.force_login(student);assignment.status="closed";assignment.save()
    assert upload(client,assignment).status_code==400

def test_submission_total_and_file_count_limits(client,student,assignment):
    client.force_login(student)
    files=[SimpleUploadedFile(f"r{i}.txt",b"x"*(18*1024*1024)) for i in range(3)]
    assert upload(client,assignment,report_files=files).status_code==413
    assert upload(client,assignment,report_files=[SimpleUploadedFile(f"r{i}.txt",b"text") for i in range(21)]).status_code==413

def test_text_encoding_and_docx_tables(tmp_path):
    from docx import Document
    from plagiarism.extraction import extract_file
    path=tmp_path/"report.txt";path.write_bytes("这是编码测试，保留中文。".encode("gb18030"))
    result=extract_file(path,"report")
    assert result.status=="ready" and "保留中文" in result.text
    path=tmp_path/"report.docx";doc=Document();doc.add_paragraph("正文测试")
    table=doc.add_table(rows=1,cols=1);table.cell(0,0).text="表格中的测试文字"
    nested=table.cell(0,0).add_table(rows=1,cols=1);nested.cell(0,0).text="嵌套表格中的文字"
    doc.save(path);result=extract_file(path,"report")
    assert "表格中的测试文字" in result.text and "嵌套表格中的文字" in result.text
    path.write_bytes(b"not a docx")
    assert extract_file(path,"report").status=="failed"

def test_draft_and_cross_course_assignment_hidden(client,student,other_teacher,assignment):
    client.force_login(other_teacher)
    assert client.get(f"/api/assignments/{assignment.id}/").status_code==404
    client.force_login(student);assignment.status="draft";assignment.save()
    assert client.get(f"/api/assignments/{assignment.id}/").status_code==404
    assignment.status="open";assignment.save();assignment.course.status="archived";assignment.course.save()
    assert upload(client,assignment).status_code==400
